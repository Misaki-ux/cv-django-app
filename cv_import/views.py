import os
import re
import json
from datetime import date
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse
from django.utils.translation import gettext_lazy as _
from django.conf import settings

from .models import ImportedCV
from accounts.models import UserProfile, Education, Experience, Skill, Language
from cv_builder.models import (
    CV, CVTemplate, CVSkill, CVLanguage, CVEducation, CVExperience,
)


def extract_text_from_pdf(file_path):
    """Extract text from PDF using pdfplumber, fallback to pytesseract for scanned PDFs."""
    import pdfplumber
    text = ""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception:
        pass

    if not text.strip():
        try:
            import pytesseract
            from PIL import Image
            from pdf2image import convert_from_path
            images = convert_from_path(file_path)
            for img in images:
                text += pytesseract.image_to_string(img) + "\n"
        except ImportError:
            try:
                import pytesseract
                from PIL import Image
                import pdfplumber as pb
                with pb.open(file_path) as pdf:
                    for page in pdf.pages:
                        img = page.to_image(resolution=300)
                        pil_img = img.original
                        text += pytesseract.image_to_string(pil_img) + "\n"
            except Exception:
                pass
    return text


def extract_text_from_docx(file_path):
    """Extract text from DOCX files."""
    from docx import Document
    doc = Document(file_path)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text


def parse_cv_text(text):
    """Extract common CV sections from PDF/DOCX text without guessing headings as names."""
    # pdfplumber emits these placeholders when the PDF font has no usable
    # Unicode mapping. Remove the placeholders before parsing visible text.
    text = re.sub(r'\(cid:\s*\d+\)', ' ', text, flags=re.IGNORECASE)
    data = {
        'full_name': '',
        'email': '',
        'phone': '',
        'address': '',
        'summary': '',
        'education': [],
        'experience': [],
        'skills': [],
        'languages': [],
    }

    # Extract email
    email_pattern = r'[\w.+-]+@[\w-]+\.[\w.-]+'
    emails = re.findall(email_pattern, text)
    if emails:
        data['email'] = emails[0]

    # Prefer phone values from an explicitly labelled contact line and French
    # national/international formats. This avoids selecting a lone digit, date,
    # or part of the postal address as the phone number.
    phone_candidates = []
    for line in text.splitlines():
        line_without_email = re.sub(email_pattern, ' ', line)
        for candidate in re.findall(r'(?<!\w)(?:\+\s?\d{1,3}[\s()./-]*)?\d[\d\s()./-]{7,}\d(?!\w)', line_without_email):
            digit_count = sum(char.isdigit() for char in candidate)
            if 9 <= digit_count <= 15:
                value = candidate.strip(' .,;')
                if re.fullmatch(r'\d{1,2}[/.]\d{1,2}[/.]\d{2,4}', value):
                    continue
                labelled = bool(re.search(r'\b(?:t[ée]l(?:ephone|éphone)?|phone|mobile|portable)\b', line, re.I))
                french_format = bool(re.match(r'\s*(?:\+\s?33|0)', value))
                phone_candidates.append((labelled, french_format, value))
    if phone_candidates:
        phone_candidates.sort(key=lambda candidate: (candidate[0], candidate[1]), reverse=True)
        data['phone'] = phone_candidates[0][2]
    address_match = re.search(
        r'\b\d{1,4}\s+[^,\n]{3,80},\s*\d{5}\s+[A-ZÀ-ÿ][\wÀ-ÿ \'-]{1,40}',
        text,
        re.IGNORECASE,
    )
    if address_match:
        data['address'] = address_match.group(0).strip()

    section_aliases = {
        'summary': r'summary|professional summary|profile|objective|about me|profil|résumé|resume|à propos',
        'education': r'education|formations?|studies|études|academic background|parcours académique',
        'experience': r'experiences?|work experience|professional experience|exp[eé]riences? professionnelles?|emploi|career history',
        'skills': r'technical skills|key skills|compétences professionnelles|compétences techniques|compétences clés|skills|compétences|savoir-faire|points forts|strengths',
        'languages': r'languages|language skills|langues',
    }
    heading_re = re.compile(
        r'^\s*(?:[-•●▪▸*]+\s*)?(' + '|'.join(
            f'(?:{alias})' for alias in section_aliases.values()
        ) + r')\s*(?::|：|-|–)?\s*(.*?)\s*$', re.IGNORECASE
    )
    sections = {key: [] for key in section_aliases}
    current_section = None
    lines = [re.sub(r'\s+', ' ', line).strip() for line in text.replace('\r', '\n').split('\n')]
    ignored_heading_re = re.compile(
        r'^\s*(?:informations? person(?:n)?elles?|personal information|contact|points forts|highlights)\s*:?\s*$',
        re.IGNORECASE,
    )

    def is_role_heading(line):
        if not line or '(cid:' in line.casefold() or line.startswith(('•', '·', '▪', '▸', '-')):
            return False
        if len(line) > 90 or len(line.split()) > 12:
            return False
        if re.match(r'^[A-ZÀ-ÖØ-Þ][A-ZÀ-ÖØ-Þ0-9 /&’\'-]{5,}(?:\s+-\s+.+)?$', line):
            return True
        prefix = re.split(r'\s+-\s+', line, maxsplit=1)[0]
        return prefix.isupper() and len(prefix) >= 8 and any(char.isalpha() for char in prefix)

    for line in lines:
        if not line:
            continue
        heading = heading_re.match(line)
        if heading:
            heading_text = heading.group(1).casefold()
            current_section = next(
                (key for key, aliases in section_aliases.items()
                 if re.fullmatch(aliases, heading_text, re.IGNORECASE)),
                None,
            )
            if current_section and heading.group(2):
                sections[current_section].append(heading.group(2))
            continue
        if ignored_heading_re.fullmatch(line):
            current_section = None
            continue
        # PDF text order can jump from the right column back into work history.
        # Keep role headings and their following lines in the experience data.
        if current_section != 'education' and is_role_heading(line):
            current_section = 'experience'
        if current_section:
            sections[current_section].append(line)

    # Some PDFs encode names with each glyph repeated; collapse those runs
    # (e.g. "YYoohhaannnn" -> "Yohann") before using ordinary name heuristics.
    duplicated_name_parts = []
    for line in lines:
        collapsed = re.sub(
            r'(.)\1*',
            lambda match: match.group(1) * max(1, (len(match.group(0)) + 1) // 2),
            line,
            flags=re.IGNORECASE,
        ).strip()
        if (len(collapsed) >= 3 and collapsed.isalpha()
                and len(collapsed) / max(len(line), 1) <= 0.7
                and not re.search(r'\b(cid|experience|formation|compétence)\b', collapsed, re.I)):
            duplicated_name_parts.append(collapsed)
    if duplicated_name_parts:
        data['full_name'] = ' '.join(duplicated_name_parts[:2])

    # Fallback: pick a plausible short name before the CV's body headings.
    headings = re.compile('|'.join(section_aliases.values()), re.IGNORECASE)
    if not data['full_name']:
        for line in lines[:12]:
            candidate = re.sub(r'^[^\w]+|[^\w.\'-]+$', '', line).strip()
            words = candidate.split()
            if (1 < len(words) <= 5 and len(candidate) <= 60
                    and not candidate.isupper()
                    and not re.search(r'[@\d]', candidate)
                    and not headings.fullmatch(candidate)
                    and not re.search(r'\b(curriculum vitae|cv|resume|résumé)\b', candidate, re.I)):
                data['full_name'] = candidate
                break

    data['summary'] = ' '.join(sections['summary']).strip()

    def split_items(section_lines):
        items = []
        for line in section_lines:
            # Split inline lists, but keep hyphens that are part of skill names.
            items.extend(part.strip(' \t•·▪▸-*') for part in re.split(r'[,;|•·▪▸]+', line))
        return [
            item for item in items
            if item and len(item) <= 100 and '@' not in item
            and not re.search(email_pattern, item)
            and not re.search(r'\b(?:mail|e-?mail).{0,12}\b(?:address|adress|adresse)\b', item, re.I)
        ]

    data['skills'] = split_items(sections['skills'])
    data['languages'] = split_items(sections['languages'])

    date_atom = r'(?:\d{1,2}[/.]\d{1,2}[/.]\d{2,4}|\d{1,2}[/.]\d{4}|(?:19|20)\d{2}|[A-Za-zÀ-ÿ.]+\s+\d{4})'
    date_pattern = re.compile(
        rf'(?P<start>{date_atom})(?:\s*(?:-|–|—|to|à|au|jusqu[’\']?à)\s*'
        rf'(?P<end>present|current|now|présent|actuel|en cours|{date_atom}))?',
        re.IGNORECASE,
    )
    education_keyword_re = re.compile(
        r'\b(?:bac(?:\s+cvpm)?|bts|master|licence|high[ -]?school|facult[eé]|'
        r'universit[eé]|lyc[eé]e|coll[eè]ge)\b',
        re.IGNORECASE,
    )

    def parse_entries(section_lines, kind):
        if kind == 'experience' and any(is_role_heading(line) for line in section_lines):
            blocks = []
            current = []
            for line in section_lines:
                cleaned = re.sub(r'^[\s•·▪▸*-]+', '', line).strip()
                if is_role_heading(cleaned) and current:
                    blocks.append(current)
                    current = []
                if cleaned:
                    current.append(cleaned)
            if current:
                blocks.append(current)

            result = []
            for block in blocks:
                position = block[0][:200]
                company = ''
                dates = []
                descriptions = []
                for line in block[1:]:
                    date_match = date_pattern.search(line)
                    if date_match:
                        dates.append(date_match)
                    if ' - ' in line and not company:
                        candidate = line.rsplit(' - ', 1)[-1].strip(' |,;–—-')
                        if candidate and '(cid:' not in candidate.casefold():
                            company = candidate[:200]
                    repeated_glyph_line = re.sub(
                        r'(.)\1*',
                        lambda match: match.group(1) * max(1, (len(match.group(0)) + 1) // 2),
                        line,
                        flags=re.IGNORECASE,
                    )
                    if ('(cid:' in line.casefold()
                            or len(repeated_glyph_line) / max(len(line), 1) <= 0.7):
                        continue
                    readable = re.sub(r'^[\s•·▪▸*-]+', '', line).strip()
                    if readable and readable != company and not date_pattern.fullmatch(readable):
                        descriptions.append(readable)
                start = dates[0].group('start') if dates else ''
                end = (dates[0].group('end') or '') if dates else ''
                result.append({
                    'position': position,
                    'company': company,
                    'start_date': start,
                    'end_date': end,
                    'description': ' '.join(descriptions),
                })
            return result

        education_year_re = re.compile(r'(?<!\d)((?:19|20)\d{2})\s*:')
        has_year = any(education_year_re.search(line) for line in section_lines)
        has_education_keyword = any(education_keyword_re.search(line) for line in section_lines)
        if kind == 'education' and (has_year or has_education_keyword):
            blocks = []
            current = []
            for line in section_lines:
                year_match = education_year_re.search(line)
                if year_match and current:
                    blocks.append(current)
                    current = []
                elif (education_keyword_re.search(line) and current
                      and not any(education_year_re.search(previous) for previous in current)):
                    blocks.append(current)
                    current = []
                if year_match:
                    prefix = line[:year_match.start()].strip(' \t•·▪▸-*')
                    if prefix:
                        sections['skills'].append(prefix)
                    current.append(f"{year_match.group(1)}: {line[year_match.end():].strip()}")
                elif education_keyword_re.search(line):
                    current.append(line)
                elif current:
                    current.append(line)
                elif line.startswith(('•', '·', '▪', '▸')):
                    # In two-column CVs, skill bullets often precede the first
                    # dated education row in pdfplumber's reading order.
                    sections['skills'].append(line)
            if current:
                blocks.append(current)
            result = []
            for block in blocks:
                year_match = re.match(r'^((?:19|20)\d{2})\s*:\s*(.*)$', block[0])
                year = year_match.group(1) if year_match else ''
                school = year_match.group(2).strip() if year_match else block[0].strip()
                school_parts = [part.strip() for part in re.split(r'\s+-\s+', school) if part.strip()]
                institution = school_parts[0] if school_parts else school
                degree = ' - '.join(school_parts[1:])
                continuation = next((line for line in block[1:] if not line.startswith(('•', '·', '▪', '▸'))), '')
                if continuation and re.search(r'\b(?:bac|bts|licence|master|dipl[oô]me|certificat|cvpm)\b', continuation, re.I):
                    degree = continuation
                    institution = school or institution
                description = ' '.join(line for line in block[1:] if line != continuation)
                result.append({
                    'degree': (degree or school)[:200],
                    'institution': institution[:200],
                    'start_date': year,
                    'end_date': '',
                    'description': description,
                })
            return result

        entries = []
        current = []
        for line in section_lines:
            cleaned = re.sub(r'^[\s•·▪▸*-]+', '', line).strip()
            if not cleaned:
                continue
            # A dated line usually starts a new position or qualification.
            if (date_pattern.search(cleaned) and current
                    and any(date_pattern.search(previous) for previous in current)):
                entries.append(current)
                current = []
            current.append(cleaned)
        if current:
            entries.append(current)

        result = []
        for block in entries:
            joined = ' '.join(block)
            date_match = date_pattern.search(joined)
            start = date_match.group('start') if date_match else ''
            end = (date_match.group('end') or '') if date_match else ''
            if date_match:
                joined = (joined[:date_match.start()] + ' ' + joined[date_match.end():]).strip(' |,;–—-')
            parts = [part.strip(' |,;–—-') for part in re.split(r'\s+[|–—-]\s+|\s+ at \s+|\s+ chez \s+', joined, maxsplit=1, flags=re.I) if part.strip(' |,;–—-')]
            if kind == 'experience':
                position = parts[0] if parts else (block[0] if block else '')
                company = parts[1] if len(parts) > 1 else ''
                description = ' '.join(block[1:]).strip()
                if date_match:
                    description = re.sub(re.escape(date_match.group(0)), '', description).strip(' |,;–—-')
                result.append({'position': position[:200], 'company': company[:200], 'start_date': start, 'end_date': end, 'description': description})
            else:
                degree = parts[0] if parts else (block[0] if block else '')
                institution = parts[1] if len(parts) > 1 else ''
                description = ' '.join(block[1:]).strip()
                if date_match:
                    description = re.sub(re.escape(date_match.group(0)), '', description).strip(' |,;–—-')
                result.append({'degree': degree[:200], 'institution': institution[:200], 'start_date': start, 'end_date': end, 'description': description})
        return result

    education_lines = sections['education']
    # Column-based PDF extraction can append the education column out of order.
    # Prefer the original FORMATIONS block and stop before the next section.
    for index, line in enumerate(lines):
        if re.fullmatch(r'formations?\s*:?', line, re.IGNORECASE):
            education_lines = []
            for candidate in lines[index + 1:]:
                next_heading = heading_re.match(candidate)
                if next_heading:
                    next_key = next(
                        (key for key, aliases in section_aliases.items()
                         if re.fullmatch(aliases, next_heading.group(1), re.IGNORECASE)),
                        None,
                    )
                    if next_key in {'experience', 'skills', 'languages', 'summary'}:
                        break
                if ignored_heading_re.fullmatch(candidate):
                    break
                education_lines.append(candidate)
            break
    education_keyword_lines = [line for line in sections['experience'] if education_keyword_re.search(line)]
    experience_lines = [line for line in sections['experience'] if not education_keyword_re.search(line)]
    education_lines.extend(education_keyword_lines)
    data['education'] = parse_entries(education_lines, 'education')
    data['experience'] = parse_entries(experience_lines, 'experience')
    # The education section may contain interleaved bullets from the CV's
    # skills column; include those visible bullets while excluding contact info.
    data['skills'] = split_items(sections['skills'])

    return data


@login_required
def import_cv(request):
    if request.method == 'POST':
        uploaded_file = request.FILES.get('cv_file')
        if not uploaded_file:
            messages.error(request, _('Please select a file to upload.'))
            return redirect('import_cv')

        file_ext = os.path.splitext(uploaded_file.name)[1].lower()
        if file_ext not in ['.pdf', '.docx']:
            messages.error(request, _('Only PDF and DOCX files are supported.'))
            return redirect('import_cv')

        imported = ImportedCV.objects.create(
            user=request.user,
            original_file=uploaded_file,
            file_name=uploaded_file.name,
            file_type=file_ext.replace('.', ''),
            status='processing'
        )

        try:
            file_path = imported.original_file.path
            if file_ext == '.pdf':
                text = extract_text_from_pdf(file_path)
            else:
                text = extract_text_from_docx(file_path)

            imported.extracted_text = text
            parsed = parse_cv_text(text)
            imported.parsed_data = parsed
            imported.status = 'completed'
            imported.save()

            messages.success(request, _('CV imported and parsed successfully!'))
            return redirect('import_review', pk=imported.pk)

        except Exception as e:
            imported.status = 'failed'
            imported.error_message = str(e)
            imported.save()
            messages.error(request, _('Error processing CV: ') + str(e))
            return redirect('import_cv')

    imports = ImportedCV.objects.filter(user=request.user)
    return render(request, 'cv_import/import.html', {'imports': imports})


@login_required
def import_review(request, pk):
    imported = get_object_or_404(ImportedCV, pk=pk, user=request.user)
    if imported.extracted_text:
        imported.parsed_data = parse_cv_text(imported.extracted_text)
        imported.save(update_fields=['parsed_data'])
    return render(request, 'cv_import/review.html', {'imported': imported})


@login_required
def import_to_profile(request, pk):
    """Import parsed CV data into user profile."""
    imported = get_object_or_404(ImportedCV, pk=pk, user=request.user)
    data = parse_cv_text(imported.extracted_text) if imported.extracted_text else imported.parsed_data
    if data != imported.parsed_data:
        imported.parsed_data = data
        imported.save(update_fields=['parsed_data'])

    if not data:
        messages.error(request, _('No parsed data available.'))
        return redirect('import_review', pk=pk)

    profile, created = UserProfile.objects.get_or_create(user=request.user)

    # Update profile fields
    if data.get('phone') and not profile.phone:
        profile.phone = data['phone']
    if data.get('address') and not profile.address:
        profile.address = data['address']
    if data.get('summary') and (
        not profile.summary or re.search(r'\(cid:\s*\d+\)', profile.summary, re.IGNORECASE)
    ):
        profile.summary = data['summary']
    profile.save()

    # Update user name, correcting the old parser's known section-heading mistake.
    user_changed = False
    if data.get('full_name'):
        parts = data['full_name'].split(' ', 1)
        existing_name = f'{request.user.first_name} {request.user.last_name}'.strip()
        old_name_is_section_or_job = (
            bool(re.fullmatch(r'exp[eé]riences? professionnelles?', existing_name, re.IGNORECASE))
            or existing_name.casefold() in {
                item.get('position', '').casefold() for item in data.get('experience', [])
            }
        )
        bad_imported_name = old_name_is_section_or_job
        if not request.user.first_name or bad_imported_name:
            request.user.first_name = parts[0]
            user_changed = True
        if len(parts) > 1 and (not request.user.last_name or bad_imported_name):
            request.user.last_name = parts[1]
            user_changed = True
    if data.get('email') and not request.user.email:
        request.user.email = data['email']
        user_changed = True
    if user_changed:
        request.user.save()

    # Import skills
    for skill_name in data.get('skills', []):
        if not Skill.objects.filter(profile=profile, name=skill_name).exists():
            Skill.objects.create(profile=profile, name=skill_name, level=3)
    for skill in profile.skills.all():
        if ('@' in skill.name or (data.get('email') and skill.name.casefold() == data['email'].casefold())
                or re.search(r'\b(?:mail|e-?mail).{0,12}\b(?:address|adress|adresse)\b', skill.name, re.I)):
            skill.delete()

    # Import languages
    for lang_name in data.get('languages', []):
        if not Language.objects.filter(profile=profile, name=lang_name).exists():
            Language.objects.create(profile=profile, name=lang_name, proficiency='B1')

    # Create a CV Builder draft from the same parsed data. Previously this
    # action only copied a few values into the profile, leaving the builder empty.
    template = CVTemplate.objects.filter(slug='modern', is_active=True).first()
    cv_title = os.path.splitext(imported.file_name)[0][:200] or _('Imported CV')
    cv_values = {
        'full_name': data.get('full_name') or request.user.get_full_name(),
        'email': data.get('email') or request.user.email,
        'phone': data.get('phone') or profile.phone,
        'address': data.get('address') or profile.address,
        'summary': data.get('summary') or profile.summary,
    }
    # Reuse a draft created by an earlier attempt so retrying an import after
    # a partial failure does not create duplicate CVs and entries.
    cv = CV.objects.filter(
        user=request.user, title=cv_title, created_at__gte=imported.created_at,
    ).order_by('-created_at').first()
    if cv is not None:
        for field, value in cv_values.items():
            setattr(cv, field, value)
        if cv.template_id is None:
            cv.template = template
        cv.save()
    else:
        cv = CV.objects.create(user=request.user, template=template, title=cv_title, **cv_values)
    for order, skill_name in enumerate(data.get('skills', [])):
        CVSkill.objects.get_or_create(cv=cv, name=skill_name, defaults={'level': 3, 'order': order})
    for skill in cv.skills.all():
        if ('@' in skill.name or (data.get('email') and skill.name.casefold() == data['email'].casefold())
                or re.search(r'\b(?:mail|e-?mail).{0,12}\b(?:address|adress|adresse)\b', skill.name, re.I)):
            skill.delete()
    for order, lang_name in enumerate(data.get('languages', [])):
        CVLanguage.objects.get_or_create(cv=cv, name=lang_name, defaults={'proficiency': 'B1', 'order': order})

    # Older parser runs occasionally classified a degree line as a job. Move
    # those records into education when reimporting the same source CV.
    education_record_re = r'\b(?:bac(?:\s+cvpm)?|cvpm|bts|licence|master|dipl[oô]me|high[ -]?school|facult[eé]|universit[eé]|lyc[eé]e|coll[eè]ge)\b'
    for old_experience in list(cv.experiences.all()):
        if re.search(education_record_re, old_experience.position, re.I):
            if not any(
                re.search(r'\b(?:bac|cvpm|bts|licence|master|dipl[oô]me)\b', item.get('degree', ''), re.I)
                for item in data.get('education', [])
            ):
                data.setdefault('education', []).append({
                    'degree': old_experience.position,
                    'institution': old_experience.company,
                    'start_date': old_experience.start_date,
                    'end_date': old_experience.end_date,
                    'description': old_experience.description,
                })
            old_experience.delete()

    def profile_date(value):
        value = (value or '').strip()
        day = 1
        month = 1
        match = re.fullmatch(r'(\d{1,2})[/.](\d{1,2})[/.](\d{2,4})', value)
        if match:
            day, month, year = int(match.group(1)), int(match.group(2)), int(match.group(3))
        else:
            match = re.fullmatch(r'(\d{1,2})[/.](\d{4})', value)
            if match:
                month, year = int(match.group(1)), int(match.group(2))
            else:
                match = re.fullmatch(r'(\d{4})', value)
                if match:
                    year = int(match.group(1))
                else:
                    match = re.fullmatch(r'([A-Za-zÀ-ÿ.]+)\s+(\d{4})', value, re.I)
                    if not match:
                        return None
                    month_name = match.group(1).casefold().rstrip('.')
                    month_names = {
                        'jan': 1, 'january': 1, 'janvier': 1,
                        'feb': 2, 'february': 2, 'février': 2, 'fevrier': 2,
                        'mar': 3, 'march': 3, 'mars': 3,
                        'apr': 4, 'april': 4, 'avril': 4,
                        'may': 5, 'mai': 5,
                        'jun': 6, 'june': 6, 'juin': 6,
                        'jul': 7, 'july': 7, 'juillet': 7,
                        'aug': 8, 'august': 8, 'août': 8, 'aout': 8,
                        'sep': 9, 'sept': 9, 'september': 9, 'septembre': 9,
                        'oct': 10, 'october': 10, 'octobre': 10,
                        'nov': 11, 'november': 11, 'novembre': 11,
                        'dec': 12, 'december': 12, 'décembre': 12, 'decembre': 12,
                    }
                    month = month_names.get(month_name)
                    if not month:
                        return None
                    year = int(match.group(2))
        if year < 100:
            year += 2000 if year <= 69 else 1900
        try:
            return date(year, month, day)
        except ValueError:
            return None

    for order, item in enumerate(data.get('education', [])):
        start_date, end_date = item.get('start_date', ''), item.get('end_date', '')
        CVEducation.objects.update_or_create(
            cv=cv,
            order=order,
            defaults={
                'institution': item.get('institution', ''),
                'degree': item.get('degree', ''),
                'start_date': start_date,
                'end_date': '' if end_date.casefold() in {'present', 'current', 'now', 'présent', 'actuel', 'en cours'} else end_date,
                'description': item.get('description', ''),
            },
        )
        profile_start = profile_date(start_date)
        # The CV Builder accepts missing dates. The profile's Education table
        # requires a start date until the optional-date migration is applied.
        if profile_start and not profile.educations.filter(
            institution=item.get('institution', ''), degree=item.get('degree', ''), start_date=profile_start
        ).exists():
            profile_end = profile_date(end_date)
            Education.objects.create(
                profile=profile,
                institution=item.get('institution', '') or item.get('degree', ''),
                degree=item.get('degree', '') or item.get('institution', ''),
                start_date=profile_start,
                end_date=profile_end,
                description=item.get('description', ''),
                order=order,
            )

    for order, item in enumerate(data.get('experience', [])):
        start_date, end_date = item.get('start_date', ''), item.get('end_date', '')
        current = end_date.casefold() in {'present', 'current', 'now', 'présent', 'actuel', 'en cours'}
        CVExperience.objects.update_or_create(
            cv=cv,
            order=order,
            defaults={
                'company': item.get('company', ''),
                'position': item.get('position', ''),
                'start_date': start_date,
                'end_date': '' if current else end_date,
                'current': current,
                'description': item.get('description', ''),
            },
        )
        profile_start = profile_date(start_date)
        if profile_start and not profile.experiences.filter(
            company=item.get('company', ''), position=item.get('position', ''), start_date=profile_start
        ).exists():
            Experience.objects.create(
                profile=profile,
                company=item.get('company', '') or item.get('position', ''),
                position=item.get('position', '') or item.get('company', ''),
                start_date=profile_start,
                end_date=None if current else profile_date(end_date),
                current=current,
                description=item.get('description', ''),
                order=order,
            )

    # A retry may contain fewer parsed rows after the extraction rules are
    # improved. Remove obsolete tail rows rather than leaving those stale jobs
    # visible alongside the corrected import.
    CVEducation.objects.filter(cv=cv, order__gte=len(data.get('education', []))).delete()
    CVExperience.objects.filter(cv=cv, order__gte=len(data.get('experience', []))).delete()

    # Remove degree entries saved as profile work experience by the old parser.
    for old_experience in profile.experiences.all():
        if re.search(education_record_re, old_experience.position, re.I):
            old_experience.delete()

    messages.success(request, _('CV data imported to your profile and CV Builder!'))
    return redirect('cv_edit', pk=cv.pk)


@login_required
def import_delete(request, pk):
    imported = get_object_or_404(ImportedCV, pk=pk, user=request.user)
    if request.method == 'POST':
        if imported.original_file:
            try:
                os.remove(imported.original_file.path)
            except OSError:
                pass
        imported.delete()
        messages.success(request, _('Imported CV deleted.'))
    return redirect('import_cv')
